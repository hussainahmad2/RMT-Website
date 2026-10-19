from docx import Document
from docx.shared import Pt, RGBColor
import os

BASE = "https://rmt-usa.com"

cases = [
    {
        "n": 1,
        "title": "Class III Implantable Cardiac Monitor — CE Mark Technical File",
        "slug": "class-iii-cardiac-monitor-ce-mark",
        "client": "European MedTech Company",
        "location": "Germany",
        "year": "2023",
        "category": "Regulatory",
        "description": (
            "Led end-to-end regulatory strategy and technical documentation for a novel "
            "Class III implantable cardiac monitor, achieving CE Mark approval within 18 months."
        ),
        "challenge": (
            "The client had a near-final device design but no regulatory documentation or QMS. "
            "They needed a complete EU MDR technical file for a novel Class III implantable "
            "cardiac monitor within an aggressive 18-month timeline."
        ),
        "solution": (
            "RMT assembled a dedicated regulatory team and implemented ISO 13485 QMS in parallel "
            "with technical file development. We authored the CER, BER, Risk Management File, "
            "and complete technical documentation package."
        ),
        "outcomes": ["CE Mark Achieved", "ISO 14971 Risk File", "Full V&V Package", "Post-Market Plan"],
        "tags": ["CE Mark", "Class III", "Implantable", "Cardiac"],
        "services": [
            ("Regulatory Compliance", "/services/regulatory-compliance"),
            ("EU MDR Compliance", "/services/regulatory-compliance/eu-mdr-compliance"),
        ],
        "money_keywords": [
            "EU MDR technical file",
            "Class III medical device CE Mark",
            "ISO 14971 risk management",
            "clinical evaluation report CER",
            "medical device regulatory compliance",
            "FDA 510(k) & EU MDR compliance",
            "implantable cardiac monitor regulatory",
        ],
    },
    {
        "n": 2,
        "title": "Wearable Drug Delivery System — Full Design & Prototyping",
        "slug": "wearable-drug-delivery-design-prototyping",
        "client": "European Pharma Company",
        "location": "UAE",
        "year": "2023",
        "category": "Medical Devices",
        "description": (
            "Designed and prototyped a wearable on-body injector for biologics administration, "
            "from initial concept through design freeze and CE marking under EU MDR."
        ),
        "challenge": (
            "Develop a miniaturised, body-worn injector for high-viscosity biologics that could "
            "be self-administered by patients — meeting both pharmaceutical and medical device "
            "regulatory requirements as a combination product."
        ),
        "solution": (
            "RMT's integrated team covered mechanical design, electronics, firmware, pharmaceutical "
            "formulation, and regulatory strategy simultaneously — reducing the typical serial "
            "development timeline by 40%."
        ),
        "outcomes": ["CE Mark Achieved", "5 Prototype Iterations", "Design History File", "EU MDR Compliance"],
        "tags": ["Wearable", "Drug Delivery", "CE Mark", "EU MDR"],
        "services": [
            ("Product Development", "/services/product-development"),
            (
                "Medical Device R&D",
                "/services/engineering-product-development/research-development-engineering",
            ),
        ],
        "money_keywords": [
            "medical device product development",
            "medical device R&D",
            "wearable drug delivery device",
            "combination product development",
            "design history file DHF",
            "medical device prototyping",
            "EU MDR product design",
        ],
    },
    {
        "n": 3,
        "title": "AI-Powered Diagnostic Imaging Platform — SaMD Development",
        "slug": "ai-diagnostic-imaging-samd",
        "client": "Diagnostics Startup",
        "location": "United Kingdom",
        "year": "2022",
        "category": "Software",
        "description": (
            "Developed an AI-powered medical image analysis platform compliant with IEC 62304 "
            "and EU MDR SaMD guidance, including full software validation lifecycle documentation."
        ),
        "challenge": (
            "A diagnostic startup had a promising AI algorithm but no pathway to CE Mark as SaMD. "
            "They needed full software development lifecycle documentation and regulatory strategy."
        ),
        "solution": (
            "RMT implemented IEC 62304-compliant development processes, validated the AI/ML model "
            "against clinical datasets, and prepared the complete SaMD technical file for EU MDR submission."
        ),
        "outcomes": ["EU MDR SaMD Compliant", "IEC 62304 Lifecycle", "ML Model Validated", "CE Technical File"],
        "tags": ["AI/ML", "SaMD", "Imaging", "IEC 62304"],
        "services": [
            ("Software & AI (SaMD)", "/services/software-ai"),
            ("Regulatory Compliance", "/services/regulatory-compliance"),
        ],
        "money_keywords": [
            "SaMD software development",
            "IEC 62304 SDLC",
            "software as a medical device",
            "AI medical device software",
            "medical software validation",
            "EU MDR SaMD",
            "diagnostic imaging AI platform",
        ],
    },
    {
        "n": 4,
        "title": "Point-of-Care Biosensor Manufacturing Scale-Up",
        "slug": "poc-biosensor-manufacturing-scale-up",
        "client": "Diagnostics Company",
        "location": "Pakistan",
        "year": "2022",
        "category": "Manufacturing",
        "description": (
            "Scaled up biosensor manufacturing from lab prototype to GMP production capability, "
            "achieving ISO 13485 certification and 50,000 unit/month output."
        ),
        "challenge": (
            "A POC diagnostics company had validated their biosensor technology at lab scale but "
            "had no pathway to GMP manufacturing. They needed to scale from 500 to 50,000 units/month "
            "within 12 months."
        ),
        "solution": (
            "RMT designed the manufacturing process, implemented ISO 13485 QMS, installed and "
            "validated production equipment, and trained 25 manufacturing staff — delivering "
            "on-time and within budget."
        ),
        "outcomes": ["ISO 13485 Certified", "50k Units/Month", "Full Traceability", "Process Validated"],
        "tags": ["Biosensor", "GMP", "Scale-Up", "ISO 13485"],
        "services": [
            ("Contract Manufacturing", "/services/contract-manufacturing"),
            ("Manufacturing Capabilities", "/services/contract-manufacturing/manufacturing-capabilities"),
        ],
        "money_keywords": [
            "ISO 13485 contract manufacturing",
            "medical device manufacturing scale-up",
            "GMP medical device production",
            "cleanroom contract manufacturing",
            "process validation medical device",
            "point of care biosensor manufacturing",
            "ISO 13485 manufacturing",
        ],
    },
    {
        "n": 5,
        "title": "Drug-Device Combination — Inhaler Product Development",
        "slug": "drug-device-inhaler-product-development",
        "client": "Pharma Company",
        "location": "UAE",
        "year": "2022",
        "category": "Pharmaceutical",
        "description": (
            "Full development of a dry powder inhaler combination product from formulation through "
            "device design, clinical studies, and regulatory submission."
        ),
        "challenge": (
            "Develop a novel dry powder inhaler for a challenging peptide API — requiring simultaneous "
            "optimisation of the formulation, device mechanics, and patient usability — for multiple "
            "regulatory markets."
        ),
        "solution": (
            "RMT's integrated pharma and medical device teams co-developed the formulation and device "
            "design in parallel, conducting formative and summative usability studies and preparing "
            "the multi-market regulatory submission."
        ),
        "outcomes": [
            "Combination Product Approved",
            "3 Market Submissions",
            "Usability Study Complete",
            "Shelf Life Validated",
        ],
        "tags": ["Pharma", "Inhaler", "Combination", "Formulation"],
        "services": [
            ("Product Development", "/services/product-development"),
            ("Biomedical Devices", "/services/bmd"),
        ],
        "money_keywords": [
            "drug-device combination product",
            "medical device product development",
            "inhaler device development",
            "biomedical device design",
            "usability engineering medical device",
            "pharmaceutical medical device R&D",
            "combination product regulatory",
        ],
    },
    {
        "n": 6,
        "title": "Surgical Robotic System — Software Validation Programme",
        "slug": "surgical-robotics-software-validation",
        "client": "Surgical Robotics Company",
        "location": "United States",
        "year": "2023",
        "category": "Software",
        "description": (
            "Comprehensive software validation programme for a Class III surgical robotic system, "
            "covering IEC 62304 lifecycle, cybersecurity, and human factors."
        ),
        "challenge": (
            "An established surgical robotics company needed a complete software V&V programme for "
            "a major software update — with FDA submission deadline in 9 months."
        ),
        "solution": (
            "RMT provided a dedicated V&V team that reviewed the software architecture, wrote and "
            "executed test protocols covering all software safety classes, and produced a complete "
            "V&V summary report for the FDA submission."
        ),
        "outcomes": [
            "V&V Package Submitted",
            "Cybersecurity Documented",
            "Human Factors Complete",
            "Zero Major Defects",
        ],
        "tags": ["Surgical", "Robotics", "Software", "V&V"],
        "services": [
            ("Software & AI (SaMD)", "/services/software-ai"),
            ("Software Quality Assurance", "/services/software-ai/software-quality-assurance"),
        ],
        "money_keywords": [
            "medical software validation",
            "IEC 62304 verification and validation",
            "software quality assurance medical device",
            "FDA software validation",
            "surgical robotics software V&V",
            "medical device cybersecurity documentation",
            "SaMD software QA",
        ],
    },
    {
        "n": 7,
        "title": "ISO 13485 QMS Implementation — Medical Device Startup",
        "slug": "iso-13485-qms-medtech-startup",
        "client": "MedTech Startup",
        "location": "Pakistan",
        "year": "2021",
        "category": "Regulatory",
        "description": (
            "Full ISO 13485 QMS implementation from scratch for a medical device startup — achieving "
            "certification in 9 months and providing the quality foundation for CE marking."
        ),
        "challenge": (
            "A startup needed ISO 13485 certification to win a key customer contract within 9 months. "
            "They had no existing quality system, no documented processes, and limited in-house QMS expertise."
        ),
        "solution": (
            "RMT implemented a lean, practical ISO 13485 QMS tailored to the startup's product scope. "
            "We wrote all procedures, trained staff, conducted internal audits, and managed the "
            "certification body interaction end-to-end."
        ),
        "outcomes": [
            "ISO 13485 Certified (9 months)",
            "Complete QMS Documentation",
            "Internal Audit Programme",
            "Customer Contract Won",
        ],
        "tags": ["ISO 13485", "QMS", "Startup", "Regulatory"],
        "services": [
            ("Regulatory Compliance", "/services/regulatory-compliance"),
            ("ISO 13485 QMS", "/services/regulatory-compliance/quality-management-system"),
        ],
        "money_keywords": [
            "ISO 13485 QMS implementation",
            "medical device quality management system",
            "ISO 13485 certification support",
            "medical device regulatory compliance",
            "QMS for medical device startup",
            "quality management system ISO 13485",
            "MedTech QMS consulting",
        ],
    },
    {
        "n": 8,
        "title": "Orthopaedic Implant — Contract Manufacturing Programme",
        "slug": "orthopaedic-implant-contract-manufacturing",
        "client": "Orthopaedic Company",
        "location": "UAE",
        "year": "2021",
        "category": "Manufacturing",
        "description": (
            "Contract manufacturing programme for a range of knee and hip orthopaedic implant "
            "components, from precision machining through sterile packaging and CE marking."
        ),
        "challenge": (
            "An orthopaedic company needed a reliable, CE-marked contract manufacturing partner for "
            "complex titanium and CoCr implant components — with full DHR documentation and sterile packaging."
        ),
        "solution": (
            "RMT established a dedicated manufacturing cell for the client, qualifying precision "
            "machining processes for titanium and CoCr, implementing sterile packaging qualification, "
            "and providing complete DHR documentation for every unit."
        ),
        "outcomes": [
            "CE Mark Maintained",
            "Full DHR per Unit",
            "Sterile Packaging Validated",
            "Zero Regulatory Issues",
        ],
        "tags": ["Orthopaedics", "Implant", "Contract Mfg", "Titanium"],
        "services": [
            ("Contract Manufacturing", "/services/contract-manufacturing"),
            ("Quality & Compliance", "/services/contract-manufacturing/quality-compliance"),
        ],
        "money_keywords": [
            "ISO 13485 contract manufacturing",
            "orthopaedic implant manufacturing",
            "medical device contract manufacturer",
            "sterile packaging validation",
            "device history record DHR",
            "implant contract manufacturing",
            "cleanroom medical manufacturing",
        ],
    },
    {
        "n": 9,
        "title": "Remote Patient Monitoring Platform — Cloud & AI Development",
        "slug": "remote-patient-monitoring-cloud-ai",
        "client": "Digital Health Company",
        "location": "United States",
        "year": "2023",
        "category": "Software",
        "description": (
            "End-to-end development of a HIPAA-compliant cloud platform for remote patient monitoring, "
            "including AI-driven alert algorithms and mobile patient application."
        ),
        "challenge": (
            "A digital health company needed to develop their RPM platform rapidly — integrating IoT "
            "medical devices, cloud data analytics, and patient/clinician mobile apps — with full "
            "HIPAA compliance and IEC 62304 documentation."
        ),
        "solution": (
            "RMT's software team built the complete platform stack — cloud infrastructure, AI alert "
            "algorithms, iOS/Android patient app, and clinical dashboard — with full IEC 62304 SDLC "
            "documentation and HIPAA security architecture."
        ),
        "outcomes": ["HIPAA Compliant", "IEC 62304 Documented", "10k Patients Onboarded", "AI Alerts Validated"],
        "tags": ["RPM", "Cloud", "AI", "HIPAA"],
        "services": [
            ("Software & AI (SaMD)", "/services/software-ai"),
            ("AI Solutions", "/services/software-ai/ai-solutions"),
        ],
        "money_keywords": [
            "remote patient monitoring software",
            "HIPAA compliant medical software",
            "SaMD cloud development",
            "AI healthcare software",
            "IEC 62304 medical app development",
            "digital health platform development",
            "medical device software AI",
        ],
    },
]

money_hub = [
    (
        "Contract Manufacturing",
        "/services/contract-manufacturing",
        "ISO 13485 contract manufacturing, cleanroom manufacturing, Class I–III production",
    ),
    (
        "Medical Device R&D",
        "/services/engineering-product-development/research-development-engineering",
        "medical device R&D, prototype engineering, design transfer",
    ),
    (
        "Product Development",
        "/services/product-development",
        "medical device product development, design controls, V&V",
    ),
    (
        "Software & AI",
        "/services/software-ai",
        "SaMD, IEC 62304, medical software, AI diagnostics",
    ),
    (
        "Regulatory Compliance",
        "/services/regulatory-compliance",
        "FDA 510(k), EU MDR, ISO 13485 QMS, technical file",
    ),
]


def set_heading_color(run, hex_color="0B3D5C"):
    run.font.color.rgb = RGBColor.from_string(hex_color)


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        set_heading_color(run, "0B3D5C" if level <= 1 else "1A5F7A")
    return h


def add_label_value(doc, label, value):
    p = doc.add_paragraph()
    r = p.add_run(f"{label}: ")
    r.bold = True
    p.add_run(value)


def add_bullet(doc, text):
    doc.add_paragraph(text, style="List Bullet")


def main():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_heading(
        "RMT Case Studies — SEO Internal Link Map & Money Keywords",
        0,
    )
    for run in title.runs:
        set_heading_color(run)

    p = doc.add_paragraph()
    p.add_run("Revive Medical Technologies Inc. (RMT)").bold = True
    doc.add_paragraph(f"Website: {BASE}")
    doc.add_paragraph(
        "Document purpose: Full inventory of live case studies, SEO wiring "
        "(Case Study → Service → Contact), and target money keywords for each page."
    )
    doc.add_paragraph("Generated for internal SEO / content / sales use.")

    add_heading(doc, "1. SEO wiring model (site-wide)", 1)
    doc.add_paragraph(
        "Every case study follows the same conversion path. This is the internal-link "
        "“map to flow” used on rmt-usa.com."
    )
    add_bullet(doc, f"Hub: {BASE}/projects")
    add_bullet(doc, "Case study page: /projects/{slug} (unique, indexable URL)")
    add_bullet(doc, "Related money service page(s): /services/…")
    add_bullet(doc, f"Lead / quote: {BASE}/contact")
    flow = doc.add_paragraph()
    flow.add_run("Flow diagram: ").bold = True
    doc.add_paragraph(
        "Case Study Page  →  Related Service (money page)  →  Contact / Request a Quote"
    )
    doc.add_paragraph(
        "Secondary links: other case studies in the same category; hub /projects; "
        "primary nav money pages in footer."
    )

    add_heading(doc, "2. Money pages (primary commercial URLs)", 1)
    for label, path, kws in money_hub:
        add_label_value(doc, label, f"{BASE}{path}")
        add_bullet(doc, f"Keywords: {kws}")

    add_heading(doc, "3. Case study inventory (9)", 1)
    doc.add_paragraph(
        "All case studies are live under /projects/{slug} and listed in sitemap.xml."
    )

    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "#"
    hdr[1].text = "Case study"
    hdr[2].text = "URL slug"
    hdr[3].text = "Category"
    hdr[4].text = "Primary money page"
    for c in cases:
        row = table.add_row().cells
        row[0].text = str(c["n"])
        row[1].text = c["title"]
        row[2].text = c["slug"]
        row[3].text = c["category"]
        row[4].text = c["services"][0][0]

    doc.add_page_break()
    add_heading(doc, "4. Full case studies with link map & keywords", 1)

    for c in cases:
        add_heading(doc, f"Case Study {c['n']}: {c['title']}", 2)
        add_label_value(doc, "Live URL", f"{BASE}/projects/{c['slug']}")
        add_label_value(doc, "Client", c["client"])
        add_label_value(doc, "Location", c["location"])
        add_label_value(doc, "Year", c["year"])
        add_label_value(doc, "Category", c["category"])
        add_label_value(doc, "Tags", ", ".join(c["tags"]))

        add_heading(doc, "Overview", 3)
        doc.add_paragraph(c["description"])

        add_heading(doc, "Challenge", 3)
        doc.add_paragraph(c["challenge"])

        add_heading(doc, "Solution", 3)
        doc.add_paragraph(c["solution"])

        add_heading(doc, "Key outcomes", 3)
        for o in c["outcomes"]:
            add_bullet(doc, o)

        add_heading(doc, "Internal link map (SEO wiring)", 3)
        doc.add_paragraph("Path visitors and crawlers should follow:")
        add_bullet(doc, f"1. Hub → {BASE}/projects")
        add_bullet(doc, f"2. Case study → {BASE}/projects/{c['slug']}")
        for i, (label, path) in enumerate(c["services"], start=3):
            add_bullet(doc, f"{i}. Related service → {label}: {BASE}{path}")
        add_bullet(
            doc,
            f"{3 + len(c['services'])}. Lead CTA → {BASE}/contact (Request a quote)",
        )

        p = doc.add_paragraph()
        p.add_run("One-line flow: ").bold = True
        svc_chain = " → ".join(f"{BASE}{path}" for _, path in c["services"])
        p.add_run(f"{BASE}/projects/{c['slug']} → {svc_chain} → {BASE}/contact")

        add_heading(doc, "Money keywords (target)", 3)
        doc.add_paragraph(
            "Primary commercial / SEO phrases this case study should support:"
        )
        for kw in c["money_keywords"]:
            add_bullet(doc, kw)

        doc.add_paragraph()

    add_heading(doc, "5. Cross-link checklist", 1)
    doc.add_paragraph("Confirm each case study page includes:")
    add_bullet(doc, "H1 = case study title")
    add_bullet(doc, "Challenge / Solution / Outcomes sections")
    add_bullet(doc, "Related services chips linking to money pages")
    add_bullet(doc, "Primary CTA: View related service")
    add_bullet(doc, "Secondary CTA: Request a quote → /contact")
    add_bullet(doc, "Sidebar related services + quote button")
    add_bullet(doc, "Links to other case studies in same category")
    add_bullet(doc, "Sitemap entry + GSC Request indexing (after deploy)")

    add_heading(doc, "6. Indexing URL list (copy for GSC)", 1)
    add_bullet(doc, f"{BASE}/projects")
    for c in cases:
        add_bullet(doc, f"{BASE}/projects/{c['slug']}")
    doc.add_paragraph()
    doc.add_paragraph("Also keep money pages indexed:")
    for label, path, _ in money_hub:
        add_bullet(doc, f"{BASE}{path}")

    out_paths = [
        r"d:\RMT-Website\docs\RMT-Case-Studies-SEO-Map.docx",
        r"c:\Users\hussa\Downloads\RMT-Case-Studies-SEO-Map.docx",
    ]
    os.makedirs(r"d:\RMT-Website\docs", exist_ok=True)
    for path in out_paths:
        doc.save(path)
        print("Saved", path)


if __name__ == "__main__":
    main()
